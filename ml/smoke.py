"""
Smoke test for the deployed ``youbank-edge-ml`` app.

    ml/.venv/bin/modal run ml/smoke.py                          # every task (cold then warm), HTTP checks, cleanup
    ml/.venv/bin/modal run ml/smoke.py --only geo.refine,docs.parse
    ml/.venv/bin/modal run ml/smoke.py --no-http                # skip the HTTP/Inngest checks
    ml/.venv/bin/modal run ml/smoke.py --report-path report.json  # also write the full JSON report

Test inputs are generated inside a Modal container (no secrets on the local machine), uploaded to R2 under
ml/test/<run>/, and every object the run writes (inputs and outputs) is deleted at the end.
The HTTP checks call the public endpoints with the bearer secret from the `youbank-edge` Modal secret,
spawn one async task and confirm its edge/ml.done event reached Inngest (HTTP 200 from inn.gs).
"""
from __future__ import annotations

import json
import threading
import time

import modal

APP = "youbank-edge-ml"
app = modal.App("youbank-edge-ml-smoke")
SECRET = modal.Secret.from_name("youbank-edge")
image = (
    modal.Image.debian_slim(python_version="3.11")
    .apt_install("espeak-ng", "fonts-dejavu-core")
    .pip_install("boto3==1.43.100", "requests", "numpy", "pillow==12.3.0", "reportlab", "python-docx==1.2.0",
                 "python-pptx==1.0.2", "openpyxl==3.1.5")
)

# Orla plant (Energy Transfer), lon -103.8965 lat 31.8107: a 2.5 km box as the app builds it (boxAround),
# the clearest Sentinel-2 L2A scenes near 2025-09-26 and 2026-09-11 (cloud < 0.01 %, tile 13SFR).
ORLA_BBOX = [-103.90971, 31.7994, -103.88329, 31.822]
ORLA_BEFORE = "S2B_MSIL2A_20250926T173009_R055_T13SFR_20250926T213049"
ORLA_AFTER = "S2B_MSIL2A_20260911T172859_R055_T13SFR_20260911T211140"


def crop_url(item: str, bbox=ORLA_BBOX, size=512) -> str:
    b = ",".join(f"{v:.5f}" for v in bbox)
    return (f"https://planetarycomputer.microsoft.com/api/data/v1/item/bbox/{b}/{size}x{size}.png"
            f"?collection=sentinel-2-l2a&item={item}&assets=visual&asset_bidx=visual%7C1%2C2%2C3&nodata=0")


GEO_INPUT = {
    "bbox": ORLA_BBOX,
    "size": 256,
    "before": {"scene": ORLA_BEFORE, "url": crop_url(ORLA_BEFORE)},
    "after": {"scene": ORLA_AFTER, "url": crop_url(ORLA_AFTER)},
    "blobs": [
        # the westernmost new pond (a connected component of pixels that turned dark)
        {"kind": 2, "x": 33.0, "y": 235.5, "bbox": [29, 224, 38, 248], "pixels": 204},
        # control: unchanged scrub/bare ground east of the highway
        {"kind": 1, "x": 220.0, "y": 200.0, "bbox": [212, 192, 228, 208], "pixels": 256},
        # both new ponds as one blob: its centroid falls on the berm between them
        {"kind": 2, "x": 39.4, "y": 235.6, "bbox": [29, 224, 51, 248], "pixels": 430},
    ],
}


# ------------------------------------------------------------------------------------------------
# Fixtures (run in Modal)
# ------------------------------------------------------------------------------------------------
def _r2():
    import os

    import boto3
    from botocore.config import Config

    return boto3.client("s3", endpoint_url=f"https://{os.environ['R2_ACCOUNT_ID']}.r2.cloudflarestorage.com",
                        aws_access_key_id=os.environ["R2_ACCESS_KEY_ID"],
                        aws_secret_access_key=os.environ["R2_SECRET_ACCESS_KEY"], region_name="auto",
                        config=Config(signature_version="s3v4", request_checksum_calculation="when_required",
                                      response_checksum_validation="when_required"))


def _font(size):
    from PIL import ImageFont

    return ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", size)


def _text_image(lines, width=1700, height=1100, size=44):
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    y = 80
    for ln in lines:
        d.text((90, y), ln, fill="black", font=_font(size))
        y += int(size * 1.6)
    return img


def _docs() -> dict[str, tuple[bytes, str]]:
    import io
    from email.message import EmailMessage

    import docx
    import openpyxl
    import requests
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    out = {}
    # 1. PDF with a text layer and a ruled table
    buf = io.BytesIO()
    st = getSampleStyleSheet()
    table = Table([["Deal", "Value ($m)", "Status"], ["Orla expansion", "450", "Announced"],
                   ["Keystone pad", "120", "Closed"], ["Midkiff cryo", "310", "Rumoured"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(buf, pagesize=letter, title="Quarterly Pipeline Review").build([
        Paragraph("Quarterly Pipeline Review", st["Title"]),
        Paragraph("Midstream operators in the Delaware Basin added processing capacity this quarter. "
                  "The table below lists the deals we are tracking and their status.", st["BodyText"]),
        Spacer(1, 12), table])
    out["text.pdf"] = (buf.getvalue(), "application/pdf")
    # 2. scanned PDF (image only, no text layer): exercises OCR and word boxes; uploaded in two parts
    scan = _text_image(["Board minutes - 14 March 2026", "The committee approved the acquisition",
                        "of the Panther plant for 275 million dollars.", "Closing is expected in the third quarter."])
    buf = io.BytesIO()
    scan.save(buf, format="PDF", resolution=200.0)
    out["scan.pdf"] = (buf.getvalue(), "application/pdf")
    # 3. an image with text
    buf = io.BytesIO()
    _text_image(["Invoice 2026-0917", "Total due: 12,400 USD", "Payable within 30 days"], 1400, 600).save(buf, "PNG")
    out["note.png"] = (buf.getvalue(), "image/png")
    # 4. DOCX: heading, paragraph, table, explicit page break, second page
    d = docx.Document()
    d.core_properties.title = "Diligence memo"
    d.add_heading("Diligence memo", 0)
    d.add_paragraph("Target: Bear Kat plant. Seller: private equity sponsor.")
    t = d.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("Metric", "Value"), ("Capacity", "200 MMcf/d"), ("Utilization", "81%")]):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    d.add_page_break()
    d.add_paragraph("Page two: risks include produced-water disposal limits.")
    buf = io.BytesIO()
    d.save(buf)
    out["memo.docx"] = (buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    # 5. PPTX: two slides, a table and speaker notes
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Permian gas processing"
    s1.placeholders[1].text = "Three new cryogenic plants in 2026"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Capacity by plant"
    tb = s2.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    for i, (a, b) in enumerate([("Plant", "MMcf/d"), ("Orla III", "275"), ("Panther", "200")]):
        tb.cell(i, 0).text, tb.cell(i, 1).text = a, b
    s2.notes_slide.notes_text_frame.text = "Mention the Orla expansion timing."
    buf = io.BytesIO()
    prs.save(buf)
    out["deck.pptx"] = (buf.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    # 6. XLSX with two sheets
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Deals"
    for r in [("Acquirer", "Target", "Value", "Date"), ("Energy Transfer", "WTG Midstream", 3.25, "2024-05-28"),
              ("ONEOK", "EnLink", 4.3, "2024-08-28"), ("Targa", "Lucid", 3.55, "2022-07-29")]:
        ws.append(r)
    wb.create_sheet("Notes").append(["Values in $bn"])
    buf = io.BytesIO()
    wb.save(buf)
    out["book.xlsx"] = (buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    # 7. CSV
    out["list.csv"] = (b"ticker,name,sector\nET,Energy Transfer,Midstream\nOKE,ONEOK,Midstream\nTRGP,Targa,Midstream\n",
                       "text/csv")
    # 8. EML with an attachment
    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = "analyst@example.com", "pm@example.com", "Orla site change"
    m["Date"] = "Tue, 29 Sep 2026 10:00:00 -0500"
    m.set_content("Satellite imagery shows two new ponds at the Orla plant since last September.")
    m.add_attachment(b"%PDF-1.4 stub", maintype="application", subtype="pdf", filename="orla-before-after.pdf")
    out["mail.eml"] = (bytes(m), "message/rfc822")
    # 9. Outlook MSG sample from the extract-msg project
    r = requests.get("https://raw.githubusercontent.com/TeamMsgExtractor/msg-extractor/master/example-msg-files/multi-to.msg",
                     timeout=60)
    r.raise_for_status()
    out["sample.msg"] = (r.content, "application/vnd.ms-outlook")
    # 10. HTML with a table
    out["page.html"] = (b"<html><head><title>Plant list</title><style>p{}</style></head><body><h1>Plants</h1>"
                        b"<p>Processing plants near Orla, Texas.</p><table><tr><th>Plant</th><th>Owner</th></tr>"
                        b"<tr><td>Orla</td><td>Energy Transfer</td></tr></table><script>var x=1;</script></body></html>",
                        "text/html")
    # 11. French plain text (language detection)
    lettre = ("Madame, Monsieur,\nNous avons le plaisir de vous informer que l'acquisition de la société "
              "a été approuvée par le conseil d'administration. La clôture est prévue au troisième trimestre.")
    out["lettre.txt"] = (lettre.encode("utf-8"), "text/plain")
    return out


def _audio() -> tuple[bytes, str, str]:
    import subprocess
    import tempfile

    import requests

    try:  # JFK inaugural address excerpt (public domain), used by the Whisper test suite
        r = requests.get("https://raw.githubusercontent.com/openai/whisper/main/tests/jfk.flac", timeout=60)
        r.raise_for_status()
        return r.content, "audio/flac", "jfk.flac (public domain)"
    except Exception:
        with tempfile.NamedTemporaryFile(suffix=".wav") as f:
            subprocess.run(["espeak-ng", "-v", "en-us", "-s", "150", "-w", f.name,
                            "The quarterly results beat expectations. Revenue grew eleven percent."], check=True)
            return open(f.name, "rb").read(), "audio/wav", "espeak-ng synthesized"


def _graph(seed: int = 7, fund: float = 3.0, director: float = 3.0, portfolio: float = 2.5, same: float = 0.7,
           fund_size: tuple = (4, 10), focused: float = 0.8, truth: bool = False) -> dict:
    """Synthetic M&A graph with planted structure that node features alone cannot see: an acquirer prefers targets
    that share a fund or a director with it, and targets in the sector where it already owns subsidiaries (most of
    a company's subsidiaries sit in one adjacent sector); serial acquirers keep acquiring (hidden appetite, visible
    only through past deals). Features carry size and a noisy sector vector, so same-sector preference and size are
    all a features-only model can use. Each deal also adds an ownership edge dated at the deal, which a leak-free
    backtest must not see. With truth=True the hidden variables come back too (for the oracle ceiling)."""
    import datetime as dt

    import numpy as np

    rng = np.random.default_rng(seed)
    n_comp, n_person, n_fund, n_firm, n_sub, n_sector, n_deals = 300, 160, 30, 40, 200, 6, 600
    sector_vec = rng.normal(size=(n_sector, 4))
    nodes, edges = [], []
    size = rng.normal(size=n_comp)
    sector = rng.integers(0, n_sector, n_comp)
    adjacent = (sector + rng.integers(1, n_sector, n_comp)) % n_sector
    appetite = rng.gamma(0.6, 1.0, n_comp)  # hidden: never a feature, only visible through past deals
    for i in range(n_comp):
        f = [size[i], rng.normal(), rng.normal(), rng.normal(), *(sector_vec[sector[i]] + 0.35 * rng.normal(size=4))]
        nodes.append({"id": i, "kind": "company", "features": [round(float(v), 4) for v in f]})
    nid = n_comp

    def add(kind, feats):
        nonlocal nid
        nodes.append({"id": nid, "kind": kind, "features": [round(float(v), 4) for v in feats]})
        nid += 1
        return nid - 1

    def old_date():
        return None if rng.random() < 0.3 else str(dt.date(2008, 1, 1) + dt.timedelta(days=int(rng.integers(0, 2900))))

    board = [set() for _ in range(n_comp)]
    funds_of = [set() for _ in range(n_comp)]
    sub_sectors = [set() for _ in range(n_comp)]
    for _ in range(n_person):
        p = add("person", [0, 0, 0, 0, *rng.normal(size=4)])
        for c in rng.choice(n_comp, int(rng.integers(2, 5)), replace=False):
            edges.append({"s": p, "d": int(c), "kind": "director", "t": old_date()})
            board[c].add(p)
    for _ in range(n_fund):
        fnode = add("fund", [rng.normal(), 0, 0, 0, *rng.normal(size=4)])
        for c in rng.choice(n_comp, int(rng.integers(*fund_size)), replace=False):
            edges.append({"s": fnode, "d": int(c), "kind": "stake", "t": old_date()})
            funds_of[c].add(fnode)
    for _ in range(n_firm):
        fm = add("firm", [rng.normal(), 0, 0, 0, *rng.normal(size=4)])
        for c in rng.choice(n_comp, int(rng.integers(5, 13)), replace=False):
            edges.append({"s": fm, "d": int(c), "kind": "advisor", "t": old_date()})
    parent_p = np.exp(size) / np.exp(size).sum()
    for _ in range(n_sub):
        parent = int(rng.choice(n_comp, p=parent_p))
        s = int(adjacent[parent]) if rng.random() < focused else int(rng.integers(0, n_sector))
        sb = add("subsidiary", [rng.normal() - 1, 0, 0, 0, *(sector_vec[s] + 0.35 * rng.normal(size=4))])
        edges.append({"s": parent, "d": sb, "kind": "subsidiary", "t": None})
        sub_sectors[parent].add(s)

    def logit_row(a):
        return np.array([fund * bool(funds_of[a] & funds_of[c]) + director * bool(board[a] & board[c])
                         + portfolio * (sector[c] in sub_sectors[a]) + same * (sector[c] == sector[a])
                         - 0.6 * (size[c] > size[a]) for c in range(n_comp)])

    L = np.stack([logit_row(a) for a in range(n_comp)])
    np.fill_diagonal(L, -np.inf)
    acq_p = np.exp(0.8 * size) * appetite
    acq_p /= acq_p.sum()
    start = dt.date(2016, 1, 1)
    dates = sorted(start + dt.timedelta(days=int(d)) for d in rng.integers(0, 3650, n_deals))
    deals, done = [], set()
    for day in dates:
        a = int(rng.choice(n_comp, p=acq_p))
        lg = L[a].copy()
        for c in [c for (x, c) in done if x == a]:
            lg[c] = -np.inf
        p = np.exp(lg - lg[np.isfinite(lg)].max())
        p[~np.isfinite(lg)] = 0
        b = int(rng.choice(n_comp, p=p / p.sum()))
        done.add((a, b))
        deals.append({"acquirer": a, "target": b, "t": str(day)})
        edges.append({"s": a, "d": b, "kind": "ownership", "t": str(day)})  # leak bait: dated at the deal
    g = {"nodes": nodes, "edges": edges, "deals": deals}
    if truth:
        g["truth"] = {"logits": L, "acq_p": acq_p}
    return g


def _oracle(g: dict, split: str) -> dict:
    """Ranking metrics if the generating probabilities were known: the ceiling for any model on this graph."""
    import numpy as np

    L, acq_p = g["truth"]["logits"], g["truth"]["acq_p"]
    n = L.shape[0]
    logZ = np.log(np.exp(np.where(np.isfinite(L), L, -np.inf)).sum(1))
    pairs = [(d["acquirer"], d["target"]) for d in g["deals"]]
    test = [(d["acquirer"], d["target"]) for d in g["deals"] if d["t"] >= split]
    by_acq, by_tgt = {}, {}
    for a, b in pairs:
        by_acq.setdefault(a, set()).add(b)
        by_tgt.setdefault(b, set()).add(a)
    ranks = {"t": [], "a": []}
    for a, b in test:
        for d, anchor, truth, others, s in (("t", a, b, by_acq[a], L[a]),
                                            ("a", b, a, by_tgt[b], np.log(acq_p) + L[:, b] - logZ)):
            mask = np.ones(n, bool)
            mask[anchor] = False
            for o in others:
                if o != truth:
                    mask[o] = False
            tv = s[truth]
            ranks[d].append(1 + float((s[mask] > tv).sum()) + 0.5 * float((s[mask] == tv).sum() - 1))

    def summ(r):
        r = np.asarray(r)
        return {"hits5": round(float((r <= 5).mean()), 4), "hits10": round(float((r <= 10).mean()), 4),
                "mrr": round(float((1 / r).mean()), 4), "n": int(r.size)}

    return {**summ(ranks["t"] + ranks["a"]), "asTarget": summ(ranks["t"]), "asAcquirer": summ(ranks["a"])}


def _table(seed: int = 11, n: int = 1200) -> list:
    import numpy as np

    rng = np.random.default_rng(seed)
    sectors, regions = ["Energy", "Tech", "Health", "Industrials", "Financials"], ["NA", "EU", "APAC"]
    rows = []
    for _ in range(n):
        s = int(rng.choice(5, p=[0.3, 0.25, 0.2, 0.15, 0.1]))
        revenue = float(np.exp(rng.normal(5 + 0.5 * s, 0.8)))
        margin = float(rng.normal(0.12, 0.04) if rng.random() < 0.7 else rng.normal(0.35, 0.05))
        rows.append({"sector": sectors[s], "region": regions[int(rng.choice(3, p=[0.5, 0.3, 0.2]))],
                     "revenue": round(revenue, 2), "margin": round(margin, 4) if rng.random() > 0.05 else None,
                     "employees": int(max(5, revenue * rng.uniform(2, 6))), "levered": bool(rng.random() < (0.6 if s == 0 else 0.3))})
    return rows


def _returns(seed: int = 5, T: int = 800) -> dict:
    import numpy as np

    rng = np.random.default_rng(seed)
    L = np.linalg.cholesky(np.array([[1, 0.6, 0.3], [0.6, 1, 0.4], [0.3, 0.4, 1]]))
    var, rows = np.full(3, 1e-4), []
    for _ in range(T):  # GARCH(1,1) volatility clustering with correlated Student-t shocks
        r = np.sqrt(var) * (L @ (rng.standard_t(5, size=3) / np.sqrt(5 / 3)))
        var = 1e-6 + 0.08 * r ** 2 + 0.9 * var
        rows.append([round(float(x), 6) for x in r])
    return {"columns": ["SPY", "XLE", "TLT"], "rows": rows}


def _vectors(seed: int = 3) -> dict:
    import numpy as np

    rng = np.random.default_rng(seed)
    centers = rng.normal(size=(6, 32)) * 3
    labels = np.repeat(np.arange(6), 60)
    V = centers[labels] + rng.normal(size=(360, 32))
    return {"ids": [f"doc-{i}" for i in range(360)], "vectors": np.round(V, 4).tolist()}


@app.function(image=image, secrets=[SECRET], timeout=900, cpu=1.0, memory=1024)
def make_fixtures(run: str) -> dict:
    import os

    s3, bucket, prefix = _r2(), os.environ["R2_BUCKET"], f"ml/test/{run}/"
    keys = []

    def put(name: str, body: bytes, ctype: str) -> str:
        k = prefix + name
        s3.put_object(Bucket=bucket, Key=k, Body=body, ContentType=ctype)
        keys.append(k)
        return k

    docs = {}
    for name, (body, mime) in _docs().items():
        if name == "scan.pdf":  # uploaded in two pieces to exercise `parts`
            half = len(body) // 2
            docs[name] = {"parts": [put(name + ".part0", body[:half], mime), put(name + ".part1", body[half:], mime)],
                          "mime": mime, "name": name, "ocr": "auto"}
        else:
            docs[name] = {"key": put(name, body, mime), "mime": mime, "name": name, "ocr": "auto"}
    audio, amime, source = _audio()
    half = len(audio) // 2
    audio_input = {"parts": [put("audio.part0", audio[:half], amime), put("audio.part1", audio[half:], amime)],
                   "mime": amime, "language": "auto"}
    g = _graph(truth=True)
    oracle = _oracle(g, "2023-07-01")
    g.pop("truth")
    graph_input = {"graphKey": put("graph.json", json.dumps(g).encode(), "application/json"),
                   "splitDate": "2023-07-01", "topK": 20}
    tab_input = {"dataKey": put("table.json", json.dumps(_table()).encode(), "application/json"),
                 "columns": [{"name": "sector", "type": "cat"}, {"name": "region", "type": "cat"},
                             {"name": "revenue", "type": "num"}, {"name": "margin", "type": "num"},
                             {"name": "employees", "type": "num"}, {"name": "levered", "type": "cat"}],
                 "n": 500, "seed": 3, "epochs": 300}
    series_input = {"dataKey": put("returns.json", json.dumps(_returns()).encode(), "application/json"),
                    "window": 60, "n": 64, "seed": 3, "steps": 1500}
    topics_input = {"vectorsKey": put("vectors.json", json.dumps(_vectors()).encode(), "application/json")}
    return {"keys": keys, "docs": docs, "audio": audio_input, "audioSource": source, "graph": graph_input,
            "graphSize": {"nodes": len(g["nodes"]), "edges": len(g["edges"]), "deals": len(g["deals"])},
            "graphOracle": oracle,
            "tabular": tab_input, "series": series_input, "topics": topics_input}


@app.function(image=image, secrets=[SECRET], timeout=1200, cpu=0.5, memory=512)
def http_checks(run: str, async_task: str, async_input: dict) -> dict:
    import os

    import requests

    api = modal.Function.from_name(APP, "api").get_web_url()
    st = modal.Function.from_name(APP, "status").get_web_url()
    auth = {"Authorization": f"Bearer {os.environ['EDGE_ML_SECRET']}"}
    out = {"apiUrl": api, "statusUrl": st}
    out["noAuth"] = requests.post(api, json={"task": "health"}, timeout=60).status_code
    out["wrongAuth"] = requests.post(api, json={"task": "health"}, headers={"Authorization": "Bearer nope"},
                                     timeout=60).status_code
    out["unknownTask"] = requests.post(api, json={"task": "nope"}, headers=auth, timeout=60).status_code
    out["badCallId"] = requests.post(st, json={"callId": "x"}, headers=auth, timeout=60).status_code
    for label in ("healthFirst", "healthSecond"):
        t0 = time.time()
        r = requests.post(api, json={"task": "health", "input": {}, "async": False, "correlation": f"smoke:{run}"},
                          headers=auth, timeout=120)
        body = r.json()
        out[label] = {"http": r.status_code, "ok": body.get("ok"), "latency": round(time.time() - t0, 2),
                      "seconds": body.get("seconds"), "costUsd": body.get("costUsd"),
                      "checks": (body.get("result") or {}).get("checks"),
                      "versions": {k: (body.get("result") or {}).get("versions", {}).get(k)
                                   for k in ("python", "modal", "torch", "faster-whisper", "torch-geometric")}}
    t0 = time.time()
    r = requests.post(api, json={"task": async_task, "input": async_input, "async": True, "correlation": f"smoke:{run}"},
                      headers=auth, timeout=60)
    spawn = r.json()
    out["asyncSpawn"] = {"http": r.status_code, **spawn, "latency": round(time.time() - t0, 2)}
    call_id = spawn.get("callId")
    polls, final = 0, None
    while call_id and time.time() - t0 < 900:
        polls += 1
        s = requests.post(st, json={"callId": call_id}, headers=auth, timeout=60).json()
        if s.get("done"):
            final = s
            break
        time.sleep(2)
    out["asyncStatus"] = {"polls": polls, "doneAfter": round(time.time() - t0, 2),
                          "done": bool(final), "ok": (final or {}).get("ok"), "error": (final or {}).get("error"),
                          "result": (final or {}).get("result"), "seconds": (final or {}).get("seconds"),
                          "costUsd": (final or {}).get("costUsd")}
    if call_id:
        env = modal.FunctionCall.from_id(call_id).get(timeout=60)
        out["asyncEnvelope"] = {"callIdMatches": env.get("callId") == call_id, "notified": env.get("notified"),
                                "ok": env.get("ok")}
    return out


@app.function(image=image, secrets=[SECRET], timeout=300, cpu=0.5, memory=1024)
def peek(keys: list) -> dict:
    """A short look at each output before cleanup deletes it."""
    import os

    s3, bucket, out = _r2(), os.environ["R2_BUCKET"], {}
    for k in keys:
        if not k.endswith(".json"):
            out[k] = {"bytes": s3.head_object(Bucket=bucket, Key=k)["ContentLength"]}
            continue
        d = json.loads(s3.get_object(Bucket=bucket, Key=k)["Body"].read())
        if k.startswith("ml/parsed/"):
            out[k] = {"meta": d["meta"], "pages": [
                {"n": p["n"], "text": p["text"][:200], "tables": [t[:3] for t in p["tables"][:2]],
                 "words": len(p.get("words") or []), "firstWord": (p.get("words") or [None])[0],
                 "notes": p.get("notes"), "width": p.get("width"), "height": p.get("height")} for p in d["pages"][:3]]}
        elif k.startswith("ml/transcripts/"):
            out[k] = {"language": d["language"], "duration": d["duration"], "segments": len(d["segments"]),
                      "text": " ".join(x["text"] for x in d["segments"])[:400]}
        elif "/tab-" in k:
            out[k] = {"quality": d["quality"], "sample": d["rows"][:3],
                      "recipe": {x: d["recipe"].get(x) for x in ("epochsRun", "batch", "pac", "modes", "trainSeconds")}}
        elif "/ts-" in k:
            out[k] = {"stats": d["stats"], "shape": [len(d["paths"]), len(d["paths"][0]), len(d["paths"][0][0])],
                      "recipe": {x: d["recipe"].get(x) for x in ("stepsRun", "finalLoss", "trainSeconds")}}
        elif k.startswith("ml/topics/"):
            out[k] = {"clusters": d["clusters"], "k": d["k"], "silhouette": d["silhouette"], "firstPoint": d["points"][0]}
        elif k.startswith("ml/predictions/"):
            out[k] = {"companies": len(d["targets"]), "sampleTargets": list(d["targets"].items())[:1],
                      "sampleAcquirers": list(d["acquirers"].items())[:1]}
    return out


@app.function(image=image, secrets=[SECRET], timeout=300, cpu=0.25, memory=512)
def cleanup(keys: list, prefix: str) -> dict:
    import os

    s3, bucket = _r2(), os.environ["R2_BUCKET"]
    listed = []
    for page in s3.get_paginator("list_objects_v2").paginate(Bucket=bucket, Prefix=prefix):
        listed += [o["Key"] for o in page.get("Contents", [])]
    todo = sorted({k for k in keys if isinstance(k, str) and k.startswith("ml/")} | set(listed))
    for i in range(0, len(todo), 1000):
        s3.delete_objects(Bucket=bucket, Delete={"Objects": [{"Key": k} for k in todo[i:i + 1000]], "Quiet": True})
    remaining = []
    for k in todo:
        try:
            s3.head_object(Bucket=bucket, Key=k)
            remaining.append(k)
        except Exception:
            pass
    return {"deleted": len(todo) - len(remaining), "remaining": remaining}


# ------------------------------------------------------------------------------------------------
# Local driver
# ------------------------------------------------------------------------------------------------
def _output_keys(result) -> list:
    if not isinstance(result, dict):
        return []
    return [v for k, v in result.items() if k in ("key", "modelKey", "predictionsKey") and isinstance(v, str)]


@app.local_entrypoint()
def main(only: str = "", no_http: bool = False, keep: bool = False, report_path: str = ""):
    run = time.strftime("%Y%m%dT%H%M%S")
    wanted = {t.strip() for t in only.split(",") if t.strip()} or None
    print(f"smoke run {run}: building fixtures")
    fx = make_fixtures.remote(run)
    created = list(fx["keys"])
    plan = [
        ("health", "health", {}, []),
        ("geo.refine", "geo_refine", GEO_INPUT, []),
        ("docs.parse", "docs_parse", fx["docs"]["text.pdf"], [v for k, v in fx["docs"].items() if k != "text.pdf"]),
        ("audio.transcribe", "audio_transcribe", fx["audio"], []),
        ("graph.train", "graph_train", fx["graph"], []),
        ("synth.tabular", "synth_tabular", fx["tabular"], []),
        ("synth.series", "synth_series", fx["series"], []),
        ("topics.map", "topics_map", fx["topics"], []),
    ]
    report, lock = {"run": run, "graphSize": fx["graphSize"], "graphOracle": fx["graphOracle"],
                    "audioSource": fx["audioSource"]}, threading.Lock()

    def call(fn, task, payload):
        t0 = time.time()
        env = fn.remote({"task": task, "input": payload, "async": False, "correlation": f"smoke:{run}"})
        return env, round(time.time() - t0, 2)

    def run_task(task, name, payload, extras):
        fn = modal.Function.from_name(APP, name)
        rec = {}
        try:
            env1, lat1 = call(fn, task, payload)
            env2, lat2 = call(fn, task, payload)
            rec = {"cold": {"latency": lat1, "seconds": env1["seconds"], "costUsd": env1["costUsd"], "ok": env1["ok"],
                            "error": env1["error"]},
                   "warm": {"latency": lat2, "seconds": env2["seconds"], "costUsd": env2["costUsd"], "ok": env2["ok"],
                            "error": env2["error"]},
                   "result": env1["result"]}
            keys = _output_keys(env1["result"]) + _output_keys(env2["result"])
            rec["more"] = []
            for p in extras:
                env, lat = call(fn, task, p)
                rec["more"].append({"name": p.get("name"), "latency": lat, "seconds": env["seconds"], "ok": env["ok"],
                                    "error": env["error"], "result": env["result"]})
                keys += _output_keys(env["result"])
        except Exception as e:
            rec = {"exception": f"{type(e).__name__}: {e}"}
            keys = []
        with lock:
            report[task] = rec
            created.extend(keys)
        print(f"  {task}: done")

    threads = [threading.Thread(target=run_task, args=p) for p in plan if not wanted or p[0] in wanted]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    if not no_http:
        print("  http checks")
        http = http_checks.remote(run, "topics.map", fx["topics"])
        report["http"] = http
        created.extend(_output_keys((http.get("asyncStatus") or {}).get("result")))

    report["outputs"] = sorted(set(created) - set(fx["keys"]))
    report["peek"] = peek.remote(report["outputs"])  # look at the outputs before they are deleted
    if not keep:
        report["cleanup"] = cleanup.remote(created, f"ml/test/{run}/")
    if report_path:
        with open(report_path, "w") as f:
            json.dump(report, f, indent=1, default=str)
        print(f"full report: {report_path}")
    print(json.dumps(report, indent=1, default=str)[:60000])
